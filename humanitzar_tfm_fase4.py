# -*- coding: utf-8 -*-
"""
Fase 4: insereix la descripció del ceptòmetre AccuPAR LP-80 a l'apartat 4.6
(Mesures de camp i validació) i afegeix la referència de Monteith (1977) a la
bibliografia, en ordre alfabètic i format APA.
"""
import copy
import docx

PATH = "/home/user/TFM_PAU/TFM_Pau_Planes_IRTA_v3.docx"
doc = docx.Document(PATH)


def set_text(p, new):
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def find(prefix, style=None):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix) and (style is None or p.style.name == style):
            return i
    raise LookupError(prefix[:60])


def insert_after(ref_par, text, template_par, style=None):
    el = copy.deepcopy(template_par._element)
    ref_par._element.addnext(el)
    new_p = docx.text.paragraph.Paragraph(el, ref_par._parent)
    set_text(new_p, text)
    if style:
        new_p.style = doc.styles[style]
    return new_p


# ---------------------------------------------------------------------------
# 1. Subapartat 4.6.1 amb la descripció de l'AccuPAR LP-80
# ---------------------------------------------------------------------------
tmpl_body = doc.paragraphs[find("Simultàniament a cada vol")]      # cos de text normal
anchor = doc.paragraphs[find("Variables meteorològiques: temperatura")]  # últim pic de la llista

blocs = [
    ("4.6.1. Mesura de la fPAR i del LAI amb el ceptòmetre AccuPAR LP-80", "Heading 3"),
    ("La radiació que impulsa la fotosíntesi és la radiació fotosintèticament activa "
     "(PAR), que ocupa la franja de l'espectre compresa entre 400 i 700 nm. Monteith "
     "(1977) va observar que la producció de matèria seca d'una coberta vegetal depèn "
     "directament de la quantitat de PAR que el dosser arriba a interceptar, la fracció "
     "que es coneix com a fPAR. Dit d'una altra manera, la fPAR indica quina proporció "
     "de la llum d'aquesta franja aprofita el dosser. Per la seva banda, l'índex d'àrea "
     "foliar (LAI, Leaf Area Index) descriu la superfície foliar total per unitat de "
     "superfície de sòl o de dosser, i està relacionat directament amb la capacitat de "
     "la planta d'interceptar radiació, fotosintetitzar i produir biomassa.", None),
    ("Sota una coberta vegetal, la radiació pot passar d'insolació plena a valors "
     "propers a zero en pocs centímetres. Per això, mesurar la fPAR amb fiabilitat "
     "obliga a prendre moltes lectures en punts i alçades diferents per sota del dosser. "
     "Totes dues variables serveixen per seguir el desenvolupament i el creixement del "
     "cultiu, ja que una fPAR alta apunta a una bona absorció de radiació per a la "
     "fotosíntesi. També permeten estimar la producció primària, és a dir, la biomassa "
     "que pot generar un cultiu o un ecosistema, i s'incorporen als models climàtics i "
     "ecològics que prediuen l'intercanvi de carboni entre la biosfera i l'atmosfera.", None),
    ("El ceptòmetre AccuPAR LP-80 de METER Group mesura la radiació PAR i la fracció "
     "que intercepta la vegetació, i a partir d'aquesta fracció estima el LAI de manera "
     "indirecta. Porta 80 sensors de PAR muntats en una barra i un sensor de PAR extern, "
     "que permet registrar alhora la radiació per sobre i per sota de la coberta. Per "
     "prendre una mesura només cal situar la barra en els punts d'interès sota el "
     "dosser: l'equip registra la radiació incident i la transmesa, i calcula "
     "automàticament la fracció interceptada i el LAI en temps real, cosa que permet "
     "comprovar sobre el terreny si els valors són coherents amb l'assaig. És un equip "
     "portàtil, alimentat amb piles, de 0,55 kg i un metre de llargada, amb memòria per "
     "a més de 4.000 dades i descàrrega per cable USB. S'utilitza per caracteritzar "
     "l'estructura de les cobertes vegetals i l'ús que fan de la PAR, per quantificar la "
     "radiació interceptada en estudis d'eficiència en l'ús de l'aigua o de competència "
     "entre cultiu i males herbes, i per alimentar models agrícoles i forestals que "
     "avaluen l'efecte de les pràctiques de maneig i dels canvis ambientals sobre la "
     "productivitat.", None),
]

ref = anchor
for text, style in blocs:
    ref = insert_after(ref, text, tmpl_body, style=style)

# ---------------------------------------------------------------------------
# 2. Referència de Monteith (1977) en ordre alfabètic
# ---------------------------------------------------------------------------
i_norman = find("Norman, J. M., Kustas, W. P., & Humes")
p_norman = doc.paragraphs[i_norman]
el = copy.deepcopy(p_norman._element)
p_norman._element.addprevious(el)
p_new = docx.text.paragraph.Paragraph(el, p_norman._parent)
set_text(p_new,
         "Monteith, J. L. (1977). Climate and the efficiency of crop production in "
         "Britain. Philosophical Transactions of the Royal Society of London. Series B, "
         "Biological Sciences, 281(980), 277–294. https://doi.org/10.1098/rstb.1977.0140")

doc.save(PATH)
print("Fase 4 completada: AccuPAR a 4.6.1 i Monteith (1977) a la bibliografia.")
