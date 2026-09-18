# -*- coding: utf-8 -*-
"""
Fase 5: crea la taula de sensors instal·lats a la parcel·la (apartat 4.3),
amb la funció de cada sensor ja redactada i les columnes tècniques buides
perquè l'autor les completi. Reajusta la numeració de les taules posteriors
i les crides que hi ha al text.
"""
import copy
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "/home/user/TFM_PAU/TFM_Pau_Planes_IRTA_v3.docx"
doc = docx.Document(PATH)
FONT = "Arial"
BUIT = ""          # cel·les per emplenar


def _font(run, size=9.5, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), FONT)


def set_text(p, new):
    runs = p.runs
    if not runs:
        p.add_run(new); return
    runs[0].text = new
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def find(prefix, style=None):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix) and (style is None or p.style.name == style):
            return i
    raise LookupError(prefix[:60])


def insert_after(ref_par, text, template_par, style=None):
    el = copy.deepcopy(template_par._element)
    ref_par._element.addnext(el)
    p = docx.text.paragraph.Paragraph(el, ref_par._parent)
    set_text(p, text)
    if style:
        p.style = doc.styles[style]
    return p


def caption_after(ref_par, label, text, template_caption):
    """Peu de taula amb camp SEQ (com els existents)."""
    el = copy.deepcopy(template_caption._element)
    ref_par._element.addnext(el)
    p = docx.text.paragraph.Paragraph(el, ref_par._parent)
    for r in p.runs:
        r._element.getparent().remove(r._element)
    r1 = p.add_run(label + " "); _font(r1, 10, bold=True)
    # camp SEQ
    r_beg = p.add_run(); _font(r_beg, 10, bold=True)
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); r_beg._element.append(fc)
    r_ins = p.add_run(); _font(r_ins, 10, bold=True)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = f' SEQ {label} \\* ARABIC '; r_ins._element.append(it)
    r_end = p.add_run(); _font(r_end, 10, bold=True)
    fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end'); r_end._element.append(fe)
    r2 = p.add_run(". " + text); _font(r2, 10)
    return p


# ---------------------------------------------------------------------------
# 1. Text d'entrada i peu de taula a l'apartat 4.3
# ---------------------------------------------------------------------------
p_intro = doc.paragraphs[find("A la parcel·la hi ha instal·lats")]
p_todo = doc.paragraphs[find("[A COMPLETAR] Model, fabricant i nombre")]
tmpl_caption = doc.paragraphs[find("Taula 2. Calendari previst", style="Caption")]

p_lead = insert_after(p_intro,
                      "La Taula 2 recull els sensors instal·lats a la parcel·la, la funció "
                      "que fa cadascun dins de l'assaig i les característiques tècniques "
                      "que cal documentar.",
                      p_intro)

p_cap = caption_after(p_lead, "Taula",
                      "Sensors instal·lats a la parcel·la experimental, funció de cadascun "
                      "i característiques tècniques.",
                      tmpl_caption)

# ---------------------------------------------------------------------------
# 2. Construcció de la taula
# ---------------------------------------------------------------------------
headers = ["Sensor o equip", "Funció dins de l'assaig", "Model i fabricant",
           "Nre.", "Tractaments i blocs", "Freqüència i sistema de registre"]
widths = [2.7, 4.6, 2.5, 0.9, 2.6, 2.7]

files = [
    ("Sensors de flux de saba (sap flow)",
     "Mesuren de manera contínua i directa la transpiració del cep. Són la "
     "referència per contrastar la transpiració que estima el model TSEB."),
    ("Cabalímetres",
     "Registren el volum d'aigua de reg aplicat a cada tractament, que es rega "
     "de manera independent. Permeten relacionar l'ETa estimada amb el consum "
     "real d'aigua de cada maneig de capçada."),
    ("Sensors d'humitat del sòl",
     "Segueixen el contingut d'aigua del sòl al llarg del perfil. Serveixen per "
     "comprovar la disponibilitat hídrica i per controlar que el reg deficitari "
     "s'aplica com estava previst."),
    ("Dendròmetres",
     "Registren les variacions del diàmetre del tronc, que són un indicador "
     "continu de l'estat hídric del cep i complementen les mesures puntuals de "
     "potencial hídric de tija."),
    ("Estació meteorològica",
     "Proporciona la temperatura de l'aire, la humitat relativa, el dèficit de "
     "pressió de vapor, la velocitat del vent i la radiació solar, que són "
     "entrades dels models TSEB i Shuttleworth–Wallace."),
    ("[Altres sensors]",
     "[Afegiu aquí qualsevol altre sensor de la parcel·la i indiqueu-ne la funció.]"),
]

table = doc.add_table(rows=1, cols=len(headers))
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(cell.paragraphs[0].add_run(h), 9.5, bold=True)
for nom, funcio in files:
    cells = table.add_row().cells
    _font(cells[0].paragraphs[0].add_run(nom), 9.5, bold=True)
    _font(cells[1].paragraphs[0].add_run(funcio), 9.5)
    for j in (2, 3, 4, 5):
        _font(cells[j].paragraphs[0].add_run(BUIT), 9.5)
for i, w in enumerate(widths):
    for row in table.rows:
        row.cells[i].width = Cm(w)

# mou la taula just després del peu
p_cap._element.addnext(table._element)

# ---------------------------------------------------------------------------
# 3. Nota per completar, després de la taula
# ---------------------------------------------------------------------------
set_text(p_todo,
         "[A COMPLETAR] Ompliu les columnes de model i fabricant, nombre d'unitats, "
         "ubicació per tractament i bloc, i freqüència i sistema de registre. Indiqueu "
         "també si les dades meteorològiques provenen d'una estació pròpia a la "
         "parcel·la o de l'estació de la Xarxa Agrometeorològica de Catalunya (XAC) més "
         "propera, tal com s'esmenta a l'apartat 4.6.")
table._element.addnext(p_todo._element)

# ---------------------------------------------------------------------------
# 4. Renumeració de les crides del text a les taules posteriors
# ---------------------------------------------------------------------------
for p in doc.paragraphs:
    if "El calendari de vols i de mesures de camp associades es resumeix a la Taula 2" in p.text:
        set_text(p, p.text.replace("es resumeix a la Taula 2", "es resumeix a la Taula 3"))
    if "La Taula 3 resumeix les entrades i les sortides" in p.text:
        set_text(p, p.text.replace("La Taula 3 resumeix", "La Taula 4 resumeix"))

# ---------------------------------------------------------------------------
# 5. Recalcula els valors desats dels camps SEQ (figures i taules)
# ---------------------------------------------------------------------------
counters = {"Figura": 0, "Taula": 0}
for p in doc.paragraphs:
    runs = p._element.findall(qn('w:r'))
    for i, r in enumerate(runs):
        instr = r.find(qn('w:instrText'))
        if instr is None or not instr.text or 'SEQ' not in instr.text:
            continue
        name = 'Figura' if 'Figura' in instr.text else ('Taula' if 'Taula' in instr.text else None)
        if name is None:
            continue
        counters[name] += 1
        # escriu el valor al node de text situat entre 'separate' i 'end'
        seen_sep = False
        for r2 in runs[i + 1:]:
            fc = r2.find(qn('w:fldChar'))
            if fc is not None:
                ft = fc.get(qn('w:fldCharType'))
                if ft == 'separate':
                    seen_sep = True; continue
                if ft == 'end':
                    if not seen_sep:      # camp sense resultat desat: n'hi afegim un
                        rpr = r.find(qn('w:rPr'))
                        r_sep = OxmlElement('w:r')
                        if rpr is not None: r_sep.append(copy.deepcopy(rpr))
                        s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate')
                        r_sep.append(s)
                        r_val = OxmlElement('w:r')
                        if rpr is not None: r_val.append(copy.deepcopy(rpr))
                        t = OxmlElement('w:t'); t.text = str(counters[name]); r_val.append(t)
                        r2.addprevious(r_sep); r2.addprevious(r_val)
                    break
            if seen_sep:
                t = r2.find(qn('w:t'))
                if t is not None:
                    t.text = str(counters[name])

doc.save(PATH)
print(f"Taula de sensors creada. Total taules: {len(doc.tables)}, "
      f"numerades {counters['Taula']}; figures {counters['Figura']}.")
