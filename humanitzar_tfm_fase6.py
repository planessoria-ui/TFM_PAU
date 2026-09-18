# -*- coding: utf-8 -*-
"""
Fase 6: afegeix l'apartat "Programa de reg i criteri de verema" (nou 4.4) amb
dues figures, renumera els apartats posteriors i actualitza les crides.
"""
import copy
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "/home/user/TFM_PAU/TFM_Pau_Planes_IRTA_v3.docx"
FIG = "/home/user/TFM_PAU/figures/"
doc = docx.Document(PATH)
FONT = "Arial"


def _font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), FONT)


def set_text(p, new):
    runs = p.runs
    if not runs:
        p.add_run(new); return
    runs[0].text = new
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def find(prefix, style=None):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix) and (style is None or p.style.name == style):
            return i
    raise LookupError(prefix[:60])


def insert_after(ref, text, template, style=None, red=False):
    el = copy.deepcopy(template._element)
    ref._element.addnext(el)
    p = docx.text.paragraph.Paragraph(el, ref._parent)
    set_text(p, text)
    if style:
        p.style = doc.styles[style]
        if style.startswith('Heading'):
            for r in p.runs:
                _font(r, 13 if style == 'Heading 2' else 12, bold=True)
    if red:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
            r.font.italic = True
            r.font.size = Pt(11)
    return p


def insert_figure(ref, filename, caption_text, source, width_cm, tmpl_caption, tmpl_body):
    # paràgraf amb la imatge
    el = copy.deepcopy(tmpl_body._element)
    ref._element.addnext(el)
    p_img = docx.text.paragraph.Paragraph(el, ref._parent)
    for r in p_img.runs:
        r._element.getparent().remove(r._element)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.add_run().add_picture(FIG + filename, width=Cm(width_cm))
    # peu amb camp SEQ
    el2 = copy.deepcopy(tmpl_caption._element)
    p_img._element.addnext(el2)
    p_cap = docx.text.paragraph.Paragraph(el2, ref._parent)
    for r in p_cap.runs:
        r._element.getparent().remove(r._element)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_cap.add_run("Figura "); _font(r1, 10, bold=True)
    rb = p_cap.add_run(); _font(rb, 10, bold=True)
    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'begin'); rb._element.append(fc)
    ri = p_cap.add_run(); _font(ri, 10, bold=True)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = ' SEQ Figura \\* ARABIC '; ri._element.append(it)
    re_ = p_cap.add_run(); _font(re_, 10, bold=True)
    fe = OxmlElement('w:fldChar'); fe.set(qn('w:fldCharType'), 'end'); re_._element.append(fe)
    r2 = p_cap.add_run(". " + caption_text + " "); _font(r2, 10)
    r3 = p_cap.add_run("Font: " + source); _font(r3, 10, italic=True)
    return p_cap


# ---------------------------------------------------------------------------
# Plantilles
# ---------------------------------------------------------------------------
tmpl_body = doc.paragraphs[find("A la parcel·la hi ha instal·lats")]
tmpl_caption = doc.paragraphs[find("Taula 1. Tractaments de maneig", style="Caption")]
anchor = doc.paragraphs[find("[A COMPLETAR] Ompliu les columnes de model")]

# ---------------------------------------------------------------------------
# Nou apartat 4.4
# ---------------------------------------------------------------------------
ref = insert_after(anchor, "4.4. Programa de reg i criteri de verema",
                   tmpl_body, style="Heading 2")

ref = insert_after(ref,
    "La dosi de reg es calcula a partir de l'evapotranspiració de referència (ETo) i del "
    "coeficient de cultiu (Kc), segons la relació ETc = ETo × Kc. L'ETo es deriva de les "
    "variables meteorològiques de la zona (temperatura de l'aire, humitat relativa, "
    "velocitat del vent i radiació solar), mentre que el Kc s'ajusta amb els valors "
    "obtinguts en campanyes anteriors de l'assaig i amb el vigor que mostra cada "
    "tractament. Tant l'aigua consumida com l'aplicada s'expressen en làmina d'aigua, on "
    "1 mm equival a 1 l/m².", tmpl_body)

ref = insert_after(ref,
    "Des del verolat, que el 2026 es va produir cap al 22 de juliol, i fins a la verema "
    "s'aplica una estratègia de reg deficitari controlat (RDC). En aquesta fase el reg es "
    "redueix fins al 25 % de la dosi que s'aplicava fins aleshores, cosa que fa baixar el "
    "potencial hídric de tija al migdia fins a valors d'entre −11 i −12 bar. El dèficit es "
    "manté dins d'aquest interval perquè el cep pugui refer reserves: la brotació i la "
    "collita de l'any següent depenen de les reserves acumulades durant la campanya "
    "actual, de manera que un dèficit més sever penalitzaria la producció de l'any vinent.",
    tmpl_body)

ref = insert_after(ref,
    "Cada tractament es rega de manera independent i la dosi es reajusta al llarg de tot "
    "el període perquè els cinc tractaments es mantinguin en el mateix potencial hídric de "
    "tija, sense llindars d'estrès diferents entre ells. Amb aquest criteri, les "
    "diferències que s'observin en l'ETa i en les variables del dosser es poden atribuir a "
    "la mida i a l'arquitectura de la capçada. La Figura 6 mostra l'evolució del potencial "
    "hídric de tija al migdia en la campanya anterior, amb el descens que provoca la "
    "imposició del dèficit a partir del verolat.", tmpl_body)

ref = insert_figure(ref, "fig6_psi_stem.png",
    "Evolució del potencial hídric de tija al migdia dels cinc tractaments de maneig de "
    "capçada durant la campanya anterior, amb el moment del verolat i les dates de verema "
    "de cada tractament.",
    "Blanco et al. (2025).", 14.5, tmpl_caption, tmpl_body)

ref = insert_after(ref,
    "La verema comença a principis de setembre i cada tractament es collita per separat "
    "quan arriba a 23,5 °Brix. Els tractaments que encara no hi han arribat es continuen "
    "regant i es deixen madurar fins que assoleixen aquest valor. El criteri segueix la "
    "pràctica del sector: si tots els tractaments es veremessin el mateix dia, només un "
    "estaria dins del rang òptim de maduració i la resta quedarien sobremadurats o poc "
    "madurs, i la comparació entre tractaments no seria vàlida. En la campanya anterior, "
    "els tractaments d'esporga severa (ES+DV i ES) van avançar la verema vuit dies "
    "respecte del maneig comercial (EM), l'esporga lleugera (EL) no la va modificar i el "
    "tractament sense esporga (SE) va endarrerir la maduresa sis dies (Blanco et al., "
    "2025).", tmpl_body)

ref = insert_after(ref,
    "El seguiment del reg i del consum es gestiona amb la plataforma Treetoscope, que "
    "integra les lectures dels sensors de flux de saba i dels cabalímetres i les expressa "
    "en làmina d'aigua. Per a cada bloc monitorat, la plataforma dona el consum real del "
    "dia anterior, la necessitat de reg del dia (en mm i en temps de reg), l'últim reg "
    "aplicat i la pluja registrada, i en representa l'evolució al costat de l'ETo i d'un "
    "indicador d'estrès (Figura 7). Aquests registres són la base per ajustar la dosi de "
    "cada tractament i, en aquest treball, la referència amb què es contrasta l'ETa "
    "estimada per teledetecció.", tmpl_body)

ref = insert_figure(ref, "fig7_treetoscope.png",
    "Interfície de la plataforma Treetoscope per a la parcel·la de vinya de l'IRTA: consum "
    "real del dia anterior, necessitat de reg del dia, últim reg aplicat i evolució del "
    "consum al costat de l'ETo, el reg, la pluja i l'indicador d'estrès.",
    "captura de pantalla de la plataforma Treetoscope.", 15.0, tmpl_caption, tmpl_body)

ref = insert_after(ref,
    "[A VERIFICAR] La Figura 6 correspon a la campanya anterior, en què la verema es va fer "
    "a 23 °Brix, mentre que el criteri d'aquesta campanya és de 23,5 °Brix. Confirmeu quin "
    "valor cal indicar a cada lloc.", tmpl_body, red=True)

ref = insert_after(ref,
    "[A COMPLETAR] Indiqueu la dosi de reg aplicada abans del verolat i el volum total "
    "aplicat per tractament durant la campanya. Comproveu també quins blocs es monitoren a "
    "Treetoscope: a la captura de la Figura 7 només n'hi apareixen tres, mentre que "
    "l'assaig té cinc tractaments.", tmpl_body, red=True)

# ---------------------------------------------------------------------------
# Renumeració dels apartats posteriors
# ---------------------------------------------------------------------------
renum = [
    ("4.4. Vols de dron i càmeres", "4.5. Vols de dron i càmeres"),
    ("4.5. Processament fotogramètric i d'imatges", "4.6. Processament fotogramètric i d'imatges"),
    ("4.6. Mesures de camp i validació", "4.7. Mesures de camp i validació"),
    ("4.6.1. Mesura de la fPAR", "4.7.1. Mesura de la fPAR i del LAI amb el ceptòmetre AccuPAR LP-80"),
    ("4.7. Models d'evapotranspiració i càlcul del CWSI", "4.8. Models d'evapotranspiració i càlcul del CWSI"),
    ("4.8. Anàlisi estadística", "4.9. Anàlisi estadística"),
]
for old, new in reversed(renum):
    for i, p in enumerate(doc.paragraphs):
        if p.style.name in ("Heading 2", "Heading 3") and p.text.strip().startswith(old):
            set_text(p, new)
            break

# ---------------------------------------------------------------------------
# Crides encreuades i estructura del treball
# ---------------------------------------------------------------------------
for p in doc.paragraphs:
    if "tal com s'esmenta a l'apartat 4.6" in p.text:
        set_text(p, p.text.replace("l'apartat 4.6", "l'apartat 4.7"))
    if p.text.strip().startswith("El treball s'organitza en vuit capítols"):
        set_text(p, p.text.replace(
            "disseny experimental, sensors instal·lats, vols de dron",
            "disseny experimental, sensors instal·lats, programa de reg, vols de dron"))

# ---------------------------------------------------------------------------
# Recalcula la numeració desada de figures i taules
# ---------------------------------------------------------------------------
counters = {"Figura": 0, "Taula": 0}
for p in doc.paragraphs:
    runs = p._element.findall(qn('w:r'))
    for i, r in enumerate(runs):
        instr = r.find(qn('w:instrText'))
        if instr is None or not instr.text or 'SEQ' not in instr.text:
            continue
        name = 'Figura' if 'Figura' in instr.text else ('Taula' if 'Taula' in instr.text else None)
        if name is None:
            continue
        counters[name] += 1
        seen_sep = False
        for r2 in runs[i + 1:]:
            fc = r2.find(qn('w:fldChar'))
            if fc is not None:
                ft = fc.get(qn('w:fldCharType'))
                if ft == 'separate':
                    seen_sep = True; continue
                if ft == 'end':
                    if not seen_sep:
                        rpr = r.find(qn('w:rPr'))
                        rs = OxmlElement('w:r')
                        if rpr is not None: rs.append(copy.deepcopy(rpr))
                        s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); rs.append(s)
                        rv = OxmlElement('w:r')
                        if rpr is not None: rv.append(copy.deepcopy(rpr))
                        t = OxmlElement('w:t'); t.text = str(counters[name]); rv.append(t)
                        r2.addprevious(rs); r2.addprevious(rv)
                    break
            if seen_sep:
                t = r2.find(qn('w:t'))
                if t is not None:
                    t.text = str(counters[name])

doc.save(PATH)
print(f"Apartat 4.4 afegit. Figures: {counters['Figura']}, taules: {counters['Taula']}.")
