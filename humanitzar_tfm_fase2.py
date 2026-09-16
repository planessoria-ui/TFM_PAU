# -*- coding: utf-8 -*-
"""
Fase 2: correccions estructurals del TFM (v3).

  - Converteix l'encapçalament buit en "4.3. Sensors instal·lats a la parcel·la"
    i hi redacta el contingut a partir de la nota de l'autor.
  - Renumera els apartats 4.3-4.7 -> 4.4-4.8.
  - Arregla la numeració de figures i taules: els camps SEQ es desen amb el
    resultat calculat, de manera que es vegin bé encara que no s'actualitzin.
  - Afegeix una nota de coherència sobre les imatges hemisfèriques (GoPro).
"""
import copy
import docx
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PATH = "/home/user/TFM_PAU/TFM_Pau_Planes_IRTA_v3.docx"
doc = docx.Document(PATH)


def set_text(p, new):
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def find(prefix, style=None):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix) and (style is None or p.style.name == style):
            return i
    raise LookupError(prefix[:60])


def insert_after(ref_par, text, template_par):
    """Insereix un paràgraf nou després de ref_par, copiant el format de template_par."""
    new_el = copy.deepcopy(template_par._element)
    ref_par._element.addnext(new_el)
    new_p = docx.text.paragraph.Paragraph(new_el, ref_par._parent)
    set_text(new_p, text)
    return new_p


# ---------------------------------------------------------------------------
# 1. Nova secció 4.3 "Sensors instal·lats a la parcel·la"
# ---------------------------------------------------------------------------
idx_note = find("EXPLICAR ELS DIEFERENTS SENSORS")
p_note = doc.paragraphs[idx_note]
p_head = doc.paragraphs[idx_note - 1]          # l'encapçalament buit
assert p_head.style.name == "Heading 2"
set_text(p_head, "4.3. Sensors instal·lats a la parcel·la")

set_text(p_note,
         "A la parcel·la hi ha instal·lats diversos sensors que prenen mesures "
         "contínues i serveixen de referència per contrastar les estimacions "
         "obtingudes per teledetecció. Els sensors de flux de saba (sap flow sensors) "
         "mesuren directament la transpiració dels ceps, i és amb aquestes dades que "
         "es compara la transpiració derivada del model TSEB. Els cabalímetres "
         "registren el volum d'aigua de reg aplicat a cada tractament, que es rega de "
         "manera independent, i permeten relacionar l'ETa estimada amb el consum real "
         "d'aigua de cada maneig de capçada.")

# Plantilla de text de marcador (vermell i cursiva), presa d'un [A CONCRETAR] existent
tmpl_ph = doc.paragraphs[find("[A CONCRETAR] Programari fotogramètric")]
insert_after(p_note,
             "[A COMPLETAR] Model, fabricant i nombre de sensors de flux de saba i de "
             "cabalímetres; en quins tractaments i blocs estan col·locats; freqüència "
             "de registre i sistema d'adquisició de dades. Afegiu-hi la resta de "
             "sensors de la parcel·la (humitat del sòl, dendròmetres, estació "
             "meteorològica) amb la funció de cadascun.",
             tmpl_ph)

# ---------------------------------------------------------------------------
# 2. Renumeració 4.3-4.7 -> 4.4-4.8
# ---------------------------------------------------------------------------
renum = [
    ("4.3. Vols de dron i sensors", "4.4. Vols de dron i càmeres"),
    ("4.4. Processament fotogramètric i d'imatges", "4.5. Processament fotogramètric i d'imatges"),
    ("4.5. Mesures de camp i validació", "4.6. Mesures de camp i validació"),
    ("4.6. Models d'evapotranspiració i càlcul del CWSI", "4.7. Models d'evapotranspiració i càlcul del CWSI"),
    ("4.7. Anàlisi estadística", "4.8. Anàlisi estadística"),
]
for old, new in reversed(renum):          # de baix a dalt per no confondre prefixos
    set_text(doc.paragraphs[find(old, style="Heading 2")], new)

# ---------------------------------------------------------------------------
# 3. Nota de coherència sobre les imatges hemisfèriques (GoPro)
# ---------------------------------------------------------------------------
idx_fipar = find("fIPAR: mesura al migdia solar")
insert_after(doc.paragraphs[find("[A CONCRETAR] Nombre de ceps mostrejats")],
             "[COHERÈNCIA A RESOLDRE] L'objectiu específic 3, la hipòtesi H4 i "
             "l'apartat 5.2 preveuen validar la fIPAR també amb imatges hemisfèriques "
             "(GoPro), però aquest apartat només descriu el ceptòmetre. Si enguany no "
             "es prenen imatges hemisfèriques, cal treure-les d'aquells tres punts; si "
             "es prenen, cal descriure-les aquí.",
             tmpl_ph)

# ---------------------------------------------------------------------------
# 4. Numeració de figures i taules: desar el resultat dels camps SEQ
# ---------------------------------------------------------------------------
counters = {"Figura": 0, "Taula": 0}
fixed = 0
for p in doc.paragraphs:
    for r in p._element.findall(qn('w:r')):
        instr = r.find(qn('w:instrText'))
        if instr is None or instr.text is None or 'SEQ' not in instr.text:
            continue
        name = 'Figura' if 'Figura' in instr.text else ('Taula' if 'Taula' in instr.text else None)
        if name is None:
            continue
        counters[name] += 1
        end = None
        for fc in r.findall(qn('w:fldChar')):
            if fc.get(qn('w:fldCharType')) == 'end':
                end = fc
        if end is None:
            continue
        # elimina qualsevol resultat antic
        for old in r.findall(qn('w:t')):
            r.remove(old)
        for fc in r.findall(qn('w:fldChar')):
            if fc.get(qn('w:fldCharType')) == 'separate':
                r.remove(fc)
        sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
        t = OxmlElement('w:t'); t.text = str(counters[name])
        end.addprevious(sep)
        end.addprevious(t)
        fixed += 1

doc.save(PATH)
print(f"Fase 2 completada. Figures: {counters['Figura']}, Taules: {counters['Taula']} "
      f"({fixed} camps numerats).")
